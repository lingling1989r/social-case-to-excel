# -*- coding: utf-8 -*-
"""QA：行高裁切量化 —— Excel 单元格行高上限 408pt，超出的行在表内显示不全。

为什么要量化而不是「感觉够长」：上一轮 35 行里有 30 行理想行高 > 408pt（最长一行
需要约 2,046pt，接近 5 屏）。如果只说「已自适应」，交付方会以为没有裁切，客户打开
看到半截文本就会怀疑数据丢了一半。正确做法是：① 把「多少行超限 / 实际行高区间 /
最长需要多少」写进「口径与说明」sheet；② 在同一个交付里给出不受行高限制的阅读入口
（独立 markdown 全文 + 把方法步骤炸成一行一步的汇总 sheet）。

本脚本按与 08_build_xlsx.py 相同的公式重算「需要多少 pt」，打印超限清单。
"""
import json, sys
from openpyxl import load_workbook

XLSX = '/tmp/social-case-to-excel-example/跨境电商案例拆解库_X_选品-黑五-关键词-Listing-站外-Reddit_20260916.xlsx'
CAP = 408  # Excel 单元格行高上限（409.5pt 保守取 408）
wb = load_workbook(XLSX)
ws = wb[wb.sheetnames[0]]
hdr = [str(c.value or '') for c in ws[1]]
# 与 build 脚本保持一致的四列宽度（按表头找列号，表头变了也能对上）
WIDTH = {'方法拆解': 84, '正文原文': 90, '正文翻译': 90, '案例研究推文': 100}
cols = [(i + 1, w) for i, h in enumerate(hdr) for k, w in WIDTH.items() if k in h]
if not cols:
    cols = [(7, 84), (21, 90), (22, 90), (23, 100)]


def lines(txt, width):
    cap = width * 1.1  # 约可容纳的拉丁字符数（中文按 2.2 计）
    out = 0
    for para in str(txt).split('\n'):
        units = sum(2.2 if ch > '\u2e80' else 1.0 for ch in para)
        out += max(1, -(-int(units) // int(cap)))
    return out


over = []
for r in range(2, ws.max_row + 1):
    need, worst = 0, 0
    for col, width in cols:
        t = ws.cell(r, column=col).value
        if t:
            h = lines(t, width) * 13.5 + 8
            if h > need:
                need, worst = h, col
    if need > CAP:
        over.append((ws.cell(r, 1).value, int(need), hdr[worst - 1], round(need / CAP, 2)))
hs = [ws.row_dimensions[r].height or 0 for r in range(2, ws.max_row + 1)]
total = ws.max_row - 1
print('需要行高 > %dpt 的行：%d / %d' % (CAP, len(over), total))
for no, need, col, ratio in over[:40]:
    print('  %s 需要 %4dpt（最长列=%-14s 约 %.2f 屏）' % (no, need, col, ratio))
print('实际行高 min/max/avg = %g / %g / %g' % (min(hs), max(hs), sum(hs) / len(hs)))
if over:
    print()
    print('→ 请把上面这组数字写进「口径与说明」sheet 的行高条目，并确认交付里有：')
    print('   ① 独立 markdown 全文   ② 一行一步的「方法拆解汇总」sheet')
