# -*- coding: utf-8 -*-
from openpyxl import load_workbook
P = '/tmp/social-case-to-excel-example/跨境电商案例拆解库_X_选品-黑五-关键词-Listing-站外-Reddit_20260916.xlsx'
wb = load_workbook(P)
ws = wb['1-案例总表']
empt = []
for r in range(2, ws.max_row + 1):
    for c in range(1, 24):
        v = ws.cell(r, c).value
        if c == 20:
            continue
        if v is None or (isinstance(v, str) and not v.strip()):
            empt.append((r, ws.cell(1, c).value))
print('空单元格:', empt)
hs = [(r, round(ws.row_dimensions[r].height or 0, 1)) for r in range(2, ws.max_row + 1)]
print('行高 min/max:', min(h for _, h in hs), max(h for _, h in hs))
print('行高样本:', hs[:6])
mx = []
for r in range(2, ws.max_row + 1):
    mx.append((len(ws.cell(r, 23).value or ''), len(ws.cell(r, 22).value or ''), len(ws.cell(r, 21).value or ''), r))
mx.sort()
print('最短推文(3):', mx[:3])
print('最长推文/译文/原文(3):', mx[-3:])
print('图片锚定类型:', {type(im).__name__ for im in ws._images})
for im in ws._images[:2]:
    print(' anchor', im.anchor._from.col, im.anchor._from.row, 'w/h', im.width, im.height)
md = open('/tmp/social-case-to-excel-example/跨境电商案例研究_42条_20260916.md', encoding='utf-8').read()
import re
nums = re.findall(r'\n## 案例 (\d+)', md)
print('md 案例号', nums[:6], '...', nums[-3:], 'count', len(nums), 'unique', len(set(nums)))
for key in ['### 故事', '### 他做对了什么', '### 你可以直接抄的动作', '### 一句话结论']:
    print(key, md.count(key))
