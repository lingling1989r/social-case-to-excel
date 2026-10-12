# -*- coding: utf-8 -*-
"""QA4：译文长度比 + sheet2 抽查。

LOW 是**启发式提示，不是失败**：阈值是 `原文 > 1200 字符 且 译文字数/原文长度 < 0.32`。
英文口语长帖（论坛求助/复盘帖，大量虚词、重复、寒暄）天然落在这个区间——实测一批 34 条里
有 10 条落在 0.28–0.31，逐条读过原文对比后**没有一条真的漏内容**，句数比基本 1:1。
所以看到 LOW 时按这个顺序判断，别直接判定翻译不合格、更不要为了过阈值灌水：
  1) 先跑句对齐检查：原文按句切开 vs 译文按「。！？」切开，句数比越接近 1 越说明没漏句；
     句数差 ≥ 3 处才需要回原文逐段补。
  2) 再看内容密度：中文字符承担的信息量高于英文，EN→ZH 的自然字符比就是 0.28–0.45；
     真正需要修的是「整句整段缺失」，不是「字符数偏少」。
  3) 把结论写进 sheet3 的「正文翻译」说明里（写明字符比区间 + 句数比口径），
     让读者知道偏短是文体差异而不是删减。
补内容时只补漏掉的句子，不要为了凑字数扩写。
"""
import json

from openpyxl import load_workbook

sel = json.load(open('/tmp/cb_sel.json', encoding='utf-8'))
rows = sel if isinstance(sel, list) else sel.get('rows', sel)
MAN = {}
for f in ['/tmp/tcb%d.py' % i for i in range(1, 7)]:
    ns = {}
    exec(open(f, encoding='utf-8').read(), ns)
    MAN.update(ns.get('MAN', {}))
print('--- 译文长度比（译文字数/原文长度）---')
lo = []
for r in rows:
    o = r.get('text', '') or ''
    t = MAN.get(r['tid'], '')
    if not t:
        continue
    ratio = len(t) / max(1, len(o))
    if len(o) > 1200 and ratio < 0.32:
        lo.append((r['tid'], len(o), len(t), round(ratio, 2), (r.get('line') or '')[:26]))
for x in lo:
    print(' LOW', x)
print('LOW count', len(lo))
wb = load_workbook('/tmp/social-case-to-excel-example/跨境电商案例拆解库_X_选品-黑五-关键词-Listing-站外-Reddit_20260916.xlsx')
ws = wb['2-方法拆解汇总']
print('--- sheet2 前 8 行 ---')
for r in range(1, 9):
    print([ (str(ws.cell(r, c).value)[:60] if ws.cell(r, c).value else None) for c in range(1, 6) ])
ws3 = wb['3-口径与说明']
print('--- sheet3 ---')
for r in range(1, ws3.max_row + 1):
    print('|', str(ws3.cell(r, 1).value)[:30], '|', str(ws3.cell(r, 2).value)[:70])
ws6 = wb['6-主题与检索口径']
print('--- sheet6 ---')
for r in range(1, ws6.max_row + 1):
    print([str(ws6.cell(r, c).value)[:34] for c in range(1, 7)])
